import io
import random
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from shipgen.dataset import build_case
from shipgen.documents import build_documents
from shipgen.generator import generate_shipment
from shipgen.render.common import LAYOUTS, SPECIMEN_NOTE, Layout, country_name, label, text
from shipgen.render.sheets import SHEETS, render_pdf
from shipgen.render.workbook import packing_list_xlsx
from shipready.documents.models import ShipmentDocuments

# Golden bundles (sea, air, "to order") and labelled cases, whose documents carry mutations:
# blank fields, a missing notify party, changed values.
BUNDLES: list[ShipmentDocuments] = [
    *(build_documents(generate_shipment(11, index)) for index in range(6)),
    *(build_case(42, index).documents for index in range(6)),
]
PORT_FIELDS = {
    "port_of_loading",
    "port_of_discharge",
    "airport_of_departure",
    "airport_of_destination",
}


def squash(value: str) -> str:
    """Text without whitespace or case, so a value wrapped over two lines still matches."""
    return re.sub(r"\s+", "", value).casefold()


def pdf_text(data: bytes) -> str:
    document = pdfium.PdfDocument(data)
    return "".join(page.get_textpage().get_text_range() for page in document)


def printed(value: Any, layout: Layout, key: str = "", parent: str = "") -> list[str]:
    """What each scalar field of a document should look like on the page."""
    if isinstance(value, BaseModel):
        return [
            shown
            for name in type(value).model_fields
            for shown in printed(getattr(value, name), layout, name, key or parent)
        ]
    if isinstance(value, tuple):
        return [shown for item in value for shown in printed(item, layout, key, parent)]
    if value is None or isinstance(value, bool) or value == "":
        return []
    if key == "country":
        # A port prints its code, which contains the country; an address prints the name.
        return [] if parent in PORT_FIELDS else [country_name(value)]
    if key.endswith("_country") or key == "country_of_origin":
        return [country_name(value)]
    return [text(value, layout)]


def documents_of(bundle: ShipmentDocuments) -> list[tuple[str, Any]]:
    return [(name, getattr(bundle, name)) for name in SHEETS if getattr(bundle, name) is not None]


@pytest.mark.parametrize("layout", LAYOUTS)
def test_every_field_is_printed(layout: Layout) -> None:
    for bundle in BUNDLES:
        for name, document in documents_of(bundle):
            page = squash(pdf_text(render_pdf([(name, document, layout)], random.Random(0))))
            missing = [v for v in printed(document, layout) if squash(v) not in page]
            assert missing == [], f"{name} ({layout}) does not show {missing}"


def test_every_page_says_it_is_a_specimen() -> None:
    bundle = BUNDLES[0]
    data = render_pdf([(name, doc, "erp") for name, doc in documents_of(bundle)], random.Random(0))
    document = pdfium.PdfDocument(data)
    assert len(document) >= len(documents_of(bundle))
    for page in document:
        assert squash(SPECIMEN_NOTE) in squash(page.get_textpage().get_text_range())


def test_several_documents_in_one_file_start_on_new_pages() -> None:
    bundle = BUNDLES[0]
    parts: list[tuple[str, Any, Layout]] = [
        ("commercial_invoice", bundle.commercial_invoice, "erp"),
        ("lut", bundle.lut, "word"),
    ]
    document = pdfium.PdfDocument(render_pdf(parts, random.Random(0)))
    first, last = (document[i].get_textpage().get_text_range() for i in (0, len(document) - 1))
    assert label("commercial_invoice", "erp") in first
    assert label("lut", "word") in last


def test_rendering_is_deterministic() -> None:
    parts: list[tuple[str, Any, Layout]] = [
        ("commercial_invoice", BUNDLES[0].commercial_invoice, "tally")
    ]
    assert render_pdf(parts, random.Random(5)) == render_pdf(parts, random.Random(5))


def test_layouts_print_values_differently() -> None:
    day, amount = date(2026, 10, 16), Decimal("157395.25")
    assert [text(day, layout) for layout in LAYOUTS] == ["16-10-2026", "16-Oct-2026", "16/10/2026"]
    assert [text(amount, layout) for layout in LAYOUTS] == [
        "157,395.25",
        "157395.25",
        "157,395.25",
    ]
    assert len({label("invoice_number", layout) for layout in LAYOUTS}) == 3


def test_packing_list_workbook_holds_every_value() -> None:
    for bundle in BUNDLES:
        packing = bundle.packing_list
        sheet = load_workbook(io.BytesIO(packing_list_xlsx(packing))).active
        assert sheet is not None
        cells = [value for row in sheet.iter_rows(values_only=True) for value in row]
        numbers = {Decimal(str(v)) for v in cells if isinstance(v, int | float)}
        dates = {v.date() for v in cells if isinstance(v, datetime)}
        texts = squash(" ".join(str(v) for v in cells if isinstance(v, str)))
        for line in packing.lines:
            assert {line.quantity, line.net_weight_kg, line.gross_weight_kg} <= numbers
            assert squash(line.description) in texts
        assert {packing.gross_weight_kg, packing.net_weight_kg, packing.volume_cbm} <= numbers
        assert packing.invoice_date in dates
        assert all(squash(c.number) in texts for c in packing.containers)


def test_packing_list_workbook_is_deterministic() -> None:
    packing = BUNDLES[0].packing_list
    assert packing_list_xlsx(packing) == packing_list_xlsx(packing)
